from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn
import re

p = Path("outputs/Water_Research完整论文草稿_结果待填.docx")
d = Document(p)
assert ZipFile(p).testzip() is None

text = "\n".join(x.text for x in d.paragraphs)
required = [
    "1. Introduction", "2. Materials and methods", "2.5 Hydro-CTGSSM",
    "2.8 Rapid phosphate-intervention experiment", "3. Results",
    "4. Discussion", "5. Conclusions", "Declarations", "References",
]
for item in required:
    assert item in text, item

assert not re.search(r"54\s*(independent|bottles|systems)", text, flags=re.I)
assert "240 min" not in text
assert "30 independent bottles" in text
assert "0, 10, 30 and 60 min" in text
assert len(d.tables) == 4

for ti, table in enumerate(d.tables, start=1):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    assert tbl_w is not None and int(tbl_w.get(qn("w:w"))) == 9360, (ti, "tblW")
    assert tbl_ind is not None and int(tbl_ind.get(qn("w:w"))) == 120, (ti, "tblInd")
    grid_widths = [int(x.get(qn("w:w"))) for x in table._tbl.tblGrid]
    assert sum(grid_widths) == 9360, (ti, grid_widths)
    for row in table.rows:
        widths = []
        for cell in row.cells:
            tcw = cell._tc.tcPr.find(qn("w:tcW"))
            widths.append(int(tcw.get(qn("w:w"))))
        assert widths == grid_widths, (ti, widths, grid_widths)

headings = [x.text for x in d.paragraphs if x.style.name.startswith("Heading")]
equations = [x.text for x in d.paragraphs if re.search(r"\([0-9]+\)$", x.text.strip())]
placeholders = [x.text for x in d.paragraphs if "[" in x.text and "]" in x.text]
words = re.findall(r"[A-Za-z0-9][A-Za-z0-9'–—\-/.]*", text)

print(f"file={p.resolve()}")
print(f"size={p.stat().st_size}")
print(f"paragraphs={len(d.paragraphs)} headings={len(headings)} tables={len(d.tables)}")
print(f"equations={len(equations)} placeholders={len(placeholders)} approx_words={len(words)}")
print("table_geometry=PASS")
print("content_structure=PASS")
