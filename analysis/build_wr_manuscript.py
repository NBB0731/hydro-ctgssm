from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = Path("outputs/Water_Research完整论文草稿_结果待填.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
LIGHT = "F4F6F9"
TABLE_FILL = "F4F6F9"
GRID = "B7C9E2"
PLACEHOLDER = "FFF2CC"
MUTED = "666666"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        el = tc_mar.find(qn(f"w:{name}"))
        if el is None:
            el = OxmlElement(f"w:{name}")
            tc_mar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def table_geometry(table, widths, indent=120):
    total = sum(widths)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(width))
        grid.append(gc)
    for row in table.rows:
        for c, width in zip(row.cells, widths):
            cell_width(c, width)


def borders(table, color=GRID, size="5"):
    tbl_pr = table._tbl.tblPr
    bd = tbl_pr.find(qn("w:tblBorders"))
    if bd is None:
        bd = OxmlElement("w:tblBorders")
        tbl_pr.append(bd)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = bd.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            bd.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_table(doc, headers, rows, widths, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0]
    repeat_header(hdr)
    for i, (cell, text) in enumerate(zip(hdr.cells, headers)):
        shade(cell, TABLE_FILL)
        cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(INK)
    for row_values in rows:
        row = table.add_row()
        no_split(row)
        for i, (cell, value) in enumerate(zip(row.cells, row_values)):
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            if alignments:
                p.alignment = alignments[i]
            p.add_run(str(value))
    table_geometry(table, widths)
    borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_numbering(doc):
    numbering = doc.part.numbering_part.element
    abs_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abs_id = max(abs_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abs_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    abstract.append(multi)
    configs = [
        (0, "decimal", "%1.", 720, 360),
        (1, "decimal", "%1.%2.", 1080, 360),
        (2, "lowerLetter", "%3)", 1440, 360),
        (3, "bullet", "•", 540, 270),
    ]
    for ilvl, fmt, text, left, hanging in configs:
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(ilvl))
        start = OxmlElement("w:start"); start.set(qn("w:val"), "1"); lvl.append(start)
        nf = OxmlElement("w:numFmt"); nf.set(qn("w:val"), fmt); lvl.append(nf)
        lt = OxmlElement("w:lvlText"); lt.set(qn("w:val"), text); lvl.append(lt)
        jc = OxmlElement("w:lvlJc"); jc.set(qn("w:val"), "left"); lvl.append(jc)
        ppr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab"); tab.set(qn("w:val"), "num"); tab.set(qn("w:pos"), str(left)); tabs.append(tab)
        ind = OxmlElement("w:ind"); ind.set(qn("w:left"), str(left)); ind.set(qn("w:hanging"), str(hanging))
        spacing = OxmlElement("w:spacing"); spacing.set(qn("w:after"), "80"); spacing.set(qn("w:line"), "290"); spacing.set(qn("w:lineRule"), "auto")
        ppr.append(tabs); ppr.append(ind); ppr.append(spacing); lvl.append(ppr)
        abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId"); ref.set(qn("w:val"), str(abs_id)); num.append(ref)
    numbering.append(num)
    return num_id


def apply_num(p, num_id, level=3):
    ppr = p._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is None:
        numpr = OxmlElement("w:numPr")
        ppr.append(numpr)
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), str(level))
    nid = OxmlElement("w:numId"); nid.set(qn("w:val"), str(num_id))
    numpr.append(ilvl); numpr.append(nid)


def bullets(doc, items, num_id):
    for item in items:
        p = doc.add_paragraph()
        apply_num(p, num_id, 3)
        p.add_run(item)


def numbered(doc, items, num_id):
    for item in items:
        p = doc.add_paragraph()
        apply_num(p, num_id, 0)
        p.add_run(item)


def add_equation(doc, eq, number):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    r = p.add_run(eq)
    r.font.name = "Cambria Math"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Cambria Math")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria Math")
    r.italic = True
    p.add_run(f"     ({number})")


def placeholder(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    r = p.add_run(text)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    r.font.highlight_color = 7  # yellow
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.492)
sec.footer_distance = Inches(0.492)

# Named override: line numbering for journal-review manuscript.
sect_pr = sec._sectPr
ln = OxmlElement("w:lnNumType")
ln.set(qn("w:countBy"), "1")
ln.set(qn("w:restart"), "newPage")
sect_pr.append(ln)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.333

for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
):
    st = styles[name]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

title_style = styles["Title"]
title_style.font.name = "Calibri"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
title_style.font.size = Pt(23)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor.from_string(INK)
title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_style.paragraph_format.space_before = Pt(0)
title_style.paragraph_format.space_after = Pt(14)
title_style.paragraph_format.keep_with_next = True

subtitle_style = styles["Subtitle"]
subtitle_style.font.name = "Calibri"
subtitle_style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
subtitle_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
subtitle_style.font.size = Pt(15)
subtitle_style.font.color.rgb = RGBColor.from_string(BLUE)
subtitle_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle_style.paragraph_format.space_before = Pt(0)
subtitle_style.paragraph_format.space_after = Pt(26)
subtitle_style.paragraph_format.keep_with_next = True

caption = styles["Caption"]
caption.font.name = "Calibri"
caption.font.size = Pt(9.5)
caption.font.italic = True
caption.font.color.rgb = RGBColor.from_string(MUTED)
caption.paragraph_format.space_before = Pt(4)
caption.paragraph_format.space_after = Pt(4)
caption.paragraph_format.keep_with_next = True

if "Equation" not in styles:
    eq_style = styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH)
    eq_style.font.name = "Cambria Math"
    eq_style.font.size = Pt(11)

num_id = add_numbering(doc)

header = sec.header
hp = header.paragraphs[0]
hp.text = "WATER RESEARCH | MANUSCRIPT DRAFT | RESULTS PENDING"
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in hp.runs:
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(MUTED)
footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fr = fp.add_run("Page ")
fr.font.size = Pt(8)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
fp._p.append(fld)

# Editorial cover
doc.add_paragraph().paragraph_format.space_after = Pt(46)
p = doc.add_paragraph("From sparse monitoring to actionable nutrient management", style="Title")
p2 = doc.add_paragraph("Hydrologically constrained probabilistic learning reveals cross-basin transferability and phosphate-intervention boundaries", style="Subtitle")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("[Author 1]ᵃ, [Author 2]ᵇ, [Author 3]ᶜ, [Corresponding author]ᵃ,*").bold = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("ᵃ [Affiliation 1]\nᵇ [Affiliation 2]\nᶜ [Affiliation 3]\n* Corresponding author: [email]")

doc.add_paragraph().paragraph_format.space_after = Pt(28)
note = doc.add_paragraph()
note.paragraph_format.left_indent = Inches(0.05)
note.paragraph_format.right_indent = Inches(0.05)
note.paragraph_format.space_before = Pt(8)
note.paragraph_format.space_after = Pt(8)
ppr = note._p.get_or_add_pPr()
shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), PLACEHOLDER); ppr.append(shd)
pbd = OxmlElement("w:pBdr")
for edge in ("top", "left", "bottom", "right"):
    el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "8"); el.set(qn("w:color"), "D6B656"); el.set(qn("w:space"), "6"); pbd.append(el)
ppr.append(pbd)
note.add_run("Draft status. ").bold = True
note.add_run(
    "All bracketed fields are intentionally left for GPU or laboratory results. They are not scientific claims and must be replaced, checked, and unhighlighted before submission."
)

doc.add_page_break()

doc.add_heading("Highlights", level=1)
bullets(doc, [
    "A continuous-time graph model handles irregular, censored river observations.",
    "Whole-basin holdouts quantify transferability beyond random data splitting.",
    "Calibrated intervals identify where water-quality inference becomes unreliable.",
    "A 60-min bottle test links model-derived chemistry boundaries to P control.",
    "Failure maps support targeted monitoring and context-specific management.",
], num_id)

doc.add_heading("Abstract", level=1)
doc.add_paragraph(
    "Sparse, irregular and spatially uneven monitoring limits the transfer of water-quality models across river basins. This study develops a hydrologically constrained continuous-time graph state-space model (Hydro-CTGSSM) that treats observations as asynchronous events, retains analytical censoring, propagates information only along directed river connections and produces calibrated prediction intervals. The model was designed for nitrate as nitrogen (NO3-N), total phosphorus (TP), dissolved oxygen (DO) and total suspended solids (TSS), with temperature, pH and electrical conductivity (EC) represented as auxiliary states. We used the UNEP GEMS/Water Global Freshwater Quality Archive v3 and linked river stations to HydroRIVERS/HydroATLAS. Evaluation separated random, forward-time and whole-basin holdouts to expose optimistic leakage from conventional data splits."
)
placeholder(doc, "[RESULTS TO INSERT: number of retained observations/stations/basins; whole-basin MAE improvement; 90% interval coverage; worst-group improvement; principal failure boundary.]")
doc.add_paragraph(
    "To connect predictive reliability with management action, GEMStat-derived pH and EC quantiles defined a rapid phosphate-removal experiment using granular ferric hydroxide (GFH). Thirty independent bottle systems covered four boundary chemistries and one central chemistry, with 0, 10, 30 and 60 min sampling."
)
placeholder(doc, "[EXPERIMENTAL RESULTS TO INSERT: GFH effect on C60/R60; pH and EC interaction; early-prediction MAE; natural-water direction check.]")
doc.add_paragraph(
    "The combined framework is intended to distinguish transferable cross-basin water-quality structure from conditions where either inference or intervention becomes unreliable. It thereby shifts machine learning from average predictive accuracy toward defensible monitoring prioritization and chemistry-specific nutrient management."
)

doc.add_paragraph().add_run("Keywords: ").bold = True
doc.paragraphs[-1].add_run("river water quality; irregular monitoring; graph neural network; censored data; conformal prediction; cross-basin transfer; phosphate; granular ferric hydroxide")

doc.add_heading("Graphical abstract concept", level=1)
doc.add_paragraph(
    "Left: irregular and censored GEMStat events mapped to a directed river network. Center: continuous-time decay, directed graph propagation and calibrated probabilistic output. Right: reliability/failure map guiding a 60-min pH × EC × GFH bottle experiment and monitoring-management decisions."
)
placeholder(doc, "[GRAPHICAL ABSTRACT ARTWORK TO BE CREATED AFTER THE MAIN RESULTS ARE AVAILABLE.]")

doc.add_heading("1. Introduction", level=1)
doc.add_paragraph(
    "River-water-quality management depends on observations that are frequently sparse in time, clustered in accessible regions and collected with heterogeneous analytical methods. These limitations are not merely a missing-data problem. They determine which processes a model can learn, which basins dominate its objective and whether apparent predictive skill survives transfer to an unmonitored basin. Global archives have improved access to standardized observations, yet their scientific value depends on models and evaluation designs that preserve irregular sampling, detection-limit information and hydrological connectivity rather than obscuring them through dense interpolation or random record splitting (Virro et al., 2021; Heinle et al., 2026)."
)
doc.add_paragraph(
    "Three gaps limit current data-driven water-quality studies. First, many sequence models presume regular time steps or fill long gaps before training. In the GEMStat records examined here, most consecutive measurements were separated by weeks to months, so daily interpolation would create many more synthetic labels than observations. Models such as GRU-D demonstrate that elapsed time and missingness patterns can be represented directly (Che et al., 2018), but a river station also exchanges information with hydrologically connected upstream reaches. Second, Euclidean proximity does not guarantee river connectivity or flow direction. Directed graph learning can represent asymmetric propagation (Li et al., 2018), while HydroRIVERS and HydroATLAS provide globally consistent network and catchment descriptors (Linke et al., 2019). Third, environmental measurements below or above reporting limits are censored, not exact values. Substitution with zero or one-half of a limit changes distributional shape and may distort extremes; likelihood-based censoring retains the information actually reported (Helsel, 2012)."
)
doc.add_paragraph(
    "A further concern is evaluation. When records from the same station or basin occur in both training and test sets, spatial attributes, monitoring practices and repeated local conditions can leak across the split. Random tests can therefore answer an easier question than the management question of interest: inference in an unmonitored basin or a future period. Whole-basin holdouts, worst-group performance and calibrated prediction intervals are needed to identify transferability rather than average fit. Distributionally robust optimization can reduce domination by data-rich groups (Sagawa et al., 2020), and conformal calibration can convert model dispersion into intervals with empirically testable coverage (Romano et al., 2019)."
)
doc.add_paragraph(
    "Reliable inference is useful only if it changes monitoring or management. We therefore added a deliberately short experiment rather than an unrelated material-development study. GEMStat pH and EC quantiles define the operating envelope, dissolved inorganic phosphorus (DIP) defines the response, and a commercially available GFH medium represents a practical intervention. Iron hydroxides are established phosphate sorbents, but water composition and pH can alter their performance (Hilbrandt et al., 2019). The experiment asks whether chemistry conditions associated with high uncertainty or elevated phosphorus risk also delimit intervention performance, and whether early observations can predict the 60-min endpoint."
)
doc.add_paragraph("The study tests four preregistered hypotheses:")
numbered(doc, [
    "A directed hydrological graph improves whole-basin transfer more than it improves random-split performance, relative to graph-free and Euclidean-graph alternatives.",
    "Continuous-time event encoding provides its largest benefit at long sampling gaps and during temporally withheld periods.",
    "Group-robust training and group-aware conformal calibration reduce worst-basin error and interval miscoverage without unacceptable loss of average accuracy.",
    "GFH reduces dissolved phosphate within 60 min, but the effect and early predictability vary across pH and EC boundaries derived independently from GEMStat.",
], num_id)
doc.add_paragraph(
    "The intended contribution is thus water-scientific rather than purely algorithmic: to quantify which cross-basin water-quality relationships transfer, where model reliability fails, and how those boundaries can inform monitoring frequency, station priority and nutrient-control decisions."
)

doc.add_heading("2. Materials and methods", level=1)
doc.add_heading("2.1 Study design", level=2)
doc.add_paragraph(
    "The workflow comprised four linked stages: (i) audit and harmonization of global river observations; (ii) station-to-network linkage and construction of static and dynamic covariates; (iii) probabilistic cross-basin learning with leakage-resistant evaluation; and (iv) a data-guided rapid experiment that tested phosphate-intervention boundaries. Model development and experimental analysis were separated: the global deep model was not trained on the bottle data, and the small experimental dataset was not used to claim validation of river-network propagation."
)

doc.add_heading("2.2 Water-quality observations", level=2)
doc.add_heading("2.2.1 GEMStat source and scope", level=3)
doc.add_paragraph(
    "We used the UNEP GEMS/Water Global Freshwater Quality Archive v3 (GFQA v3; DOI: 10.5281/zenodo.18459694), released on 2 February 2026. The downloaded archive had an MD5 checksum of 00F3EA19CE529753977EB3AEB08FBC47, matching the repository record. Version 3 contains more than 50 million measurements for 622 parameters at 22,982 stations in 42 countries from 1906 to 2024. The present analysis used river stations only and focused on records from 2010 onward to reduce historical method heterogeneity and align the model with contemporary monitoring."
)
doc.add_paragraph(
    "Primary targets were NO3-N, TP, DO and TSS. Temperature, pH and EC were auxiliary states because they were extensively observed and co-occurred with primary targets. DIP was retained for experimental design and supporting analyses but was not substituted for TP. Concentrations reported on different chemical bases (e.g., nitrate as N versus nitrate ion; phosphate as P versus phosphate ion) were converted only with explicit stoichiometric rules."
)

doc.add_heading("2.2.2 Quality control, duplicates and censoring", level=3)
doc.add_paragraph(
    "The primary analysis retained surface samples (reported depth ≤1 m or missing depth) rated Fair or Good. Unknown and Pending review records were reserved for sensitivity analysis, whereas Suspect records were excluded. Station-parameter-day duplicates were grouped by method, depth and censoring status before aggregation. Exact values were not averaged directly with detection limits. All transformation and scaling parameters were estimated from training basins only."
)
doc.add_paragraph(
    "Value flags < and > were represented as left and right censoring. For non-negative, right-skewed concentration targets, a log1p transformation was estimated within the training set; pH was not log-transformed. The transformed detection/reporting limit, rather than an arbitrary substituted concentration, entered the censored likelihood. Physically impossible values were flagged using parameter-specific bounds and station-level distributions; no global three-standard-deviation deletion rule was applied."
)

doc.add_heading("2.3 River network and environmental covariates", level=2)
doc.add_paragraph(
    "Stations were linked to HydroRIVERS reaches using coordinates and a maximum matching distance selected before outcome evaluation. Candidate matches were checked against water-body type, upstream area and named basin. Matching quality was summarized by match rate, distance distribution and a stratified manual map audit. If the reliable match rate fell below 90%, the preregistered fallback was a HydroBASINS sub-basin graph; nearest-neighbor geometry would not be described as a true river network."
)
doc.add_paragraph(
    "Directed edges followed upstream-to-downstream connectivity. Edge attributes included reach length, slope, stream order, long-term discharge proxy and an estimated travel-time proxy. Static node attributes were selected from HydroATLAS hydrology, physiography, climate, land cover/use, soils/geology and anthropogenic-pressure categories. Dynamic covariates were aggregated to the event scale from available hydroclimatic forcing [FINAL SOURCE AND VARIABLES TO CONFIRM]. Covariate definitions, units and missingness were frozen in a data dictionary before model comparison."
)

doc.add_heading("2.4 Prediction tasks and leakage-resistant splits", level=2)
add_table(doc,
    ["Task", "Test construction", "Scientific question", "Primary metrics"],
    [
        ("Random record split", "Records sampled after station grouping safeguards", "How optimistic is a conventional split?", "MAE, RMSE, R²"),
        ("Forward-time transfer", "Most recent period withheld by station/basin", "Does the model transfer to future monitoring?", "MAE, CRPS, coverage"),
        ("Whole-basin transfer", "Complete HydroBASINS units withheld", "Can the model infer unmonitored basins?", "MAE, worst-group MAE, coverage"),
        ("Masked monitoring gaps", "Contiguous 30/90/180-day blocks hidden", "How does reliability degrade with gap length?", "MAE by gap, interval width"),
        ("Extreme-state detection", "Station-season quantile events", "Are high nutrient/TSS or low DO events recovered?", "AUPRC, event recall"),
    ],
    [1700, 2700, 2860, 2100],
)
doc.add_paragraph(
    "The whole-basin split was the primary test. Basins were assigned to training, calibration and test partitions before standardization and self-supervised masking. No target values from test stations were used in pretraining, feature selection or calibration. Random splitting was retained only to quantify optimistic bias, not as the headline result."
)

doc.add_heading("2.5 Hydro-CTGSSM", level=2)
doc.add_heading("2.5.1 Event representation", level=3)
doc.add_paragraph(
    "Each observation event at station i and event index k was represented by target values vᵢₖ, an observed mask mᵢₖ, a censoring code cᵢₖ ∈ {−1,0,+1}, elapsed time Δtᵢₖ, dynamic covariates dᵢₖ and static station/basin covariates sᵢ. Multiple parameters observed on the same date were represented in a shared event vector; missing parameters retained zero-filled values only after multiplication by their mask."
)
add_equation(doc, "xᵢₖ = [vᵢₖ ⊙ mᵢₖ, mᵢₖ, cᵢₖ, dᵢₖ, sᵢ, log(1 + Δtᵢₖ)]", 1)

doc.add_heading("2.5.2 Continuous-time decay and event update", level=3)
doc.add_paragraph(
    "Between events, the previous hidden state decayed as a learned, non-negative function of elapsed time. This preserves the actual sampling interval without constructing a dense daily target series. The implemented decay and GRU update were:"
)
add_equation(doc, "γᵢₖ = softplus(Wγ log(1 + Δtᵢₖ) + bγ)", 2)
add_equation(doc, "h⁻ᵢₖ = hᵢ,ₖ₋₁ ⊙ exp[−γᵢₖ √(Δtᵢₖ)]", 3)
add_equation(doc, "hᵢₖ = GRU(SiLU[LN(Wₓxᵢₖ + bₓ)], h⁻ᵢₖ)", 4)
doc.add_paragraph(
    "Padded events did not update the hidden state. The square-root time factor follows the implemented stabilization of long intervals; alternative decay parameterizations were reserved for sensitivity analysis."
)

doc.add_heading("2.5.3 Directed hydrological message passing", level=3)
doc.add_paragraph(
    "For every directed edge u→i, an edge network converted reach attributes aᵤᵢ into a feature-wise gate. Only upstream source states contributed to downstream aggregation."
)
add_equation(doc, "gᵤᵢ = sigmoid[W₂ SiLU(W₁aᵤᵢ + b₁) + b₂]", 5)
add_equation(doc, "mᵤᵢ = (Wₘhᵤ) ⊙ gᵤᵢ;     m̄ᵢ = |N⁻(i)|⁻¹ Σᵤ∈N⁻(i) mᵤᵢ", 6)
add_equation(doc, "h′ᵢ = LN{GRU[Dropout(m̄ᵢ), hᵢ] + hᵢ}", 7)
doc.add_paragraph(
    "Three graph layers were used in the initial configuration, with hidden dimension 128 and dropout 0.1. Edge direction was encoded by construction; no undocumented mass-conservation penalty was imposed because concentrations are not loads and synchronous discharge was incomplete."
)

doc.add_heading("2.5.4 Probabilistic output and censored likelihood", level=3)
doc.add_paragraph(
    "A parameter-specific output was decoded from the final graph state concatenated with static attributes. The implemented head produced a Gaussian location μ and positive scale σ."
)
add_equation(doc, "[μᵢ, ρᵢ] = MLP([h′ᵢ, sᵢ]);     σᵢ = softplus(ρᵢ) + 10⁻⁴", 8)
doc.add_paragraph("For transformed observation y with reporting limit L, the contribution to the negative log-likelihood was:")
add_equation(doc, "ℓ = −log φ((y−μ)/σ) + log σ,                         exact", 9)
add_equation(doc, "ℓ = −log Φ((L−μ)/σ),                               left-censored", 10)
add_equation(doc, "ℓ = −log[1−Φ((L−μ)/σ)],                            right-censored", 11)
doc.add_paragraph(
    "Target weights were allowed to prevent densely measured auxiliary variables from dominating the primary endpoints. A Student-t head was not claimed in the present implementation and may be evaluated only as a preregistered sensitivity analysis."
)

doc.add_heading("2.5.5 Group-robust objective", level=3)
doc.add_paragraph(
    "Training groups were defined by major basin, with country or hydroclimatic region used when basin groups were too small. Group losses Lg were aggregated with exponentiated-gradient weights qg:"
)
add_equation(doc, "qg ← qg exp(ηLg) / Σⱼ qⱼ exp(ηLⱼ);     LDRO = Σg qg Lg", 12)
doc.add_paragraph(
    "The primary objective combined censored observation loss and GroupDRO aggregation. Structured masked-event reconstruction used only training basins. Hyperparameters were selected on a basin-disjoint validation set."
)

doc.add_heading("2.5.6 Group-aware conformal calibration", level=3)
doc.add_paragraph(
    "The validation/calibration set produced standardized nonconformity scores r = |y−μ|/max(σ,10⁻⁸). For each group with at least 50 calibration observations, the finite-sample higher quantile at level ceil[(n+1)(1−α)]/n was estimated; smaller groups used the global quantile. The 90% interval for a new event was:"
)
add_equation(doc, "I₀.₉(x) = [μ(x) − q̂g σ(x), μ(x) + q̂g σ(x)]", 13)
doc.add_paragraph(
    "Coverage was reported overall and by basin, climate zone, basin size, sampling-gap class and concentration regime. Because groupwise finite-sample guarantees require exchangeability within the calibration design, the manuscript reports empirical conditional coverage rather than claiming universal conditional validity."
)

doc.add_heading("2.6 Baselines, ablations and training", level=2)
doc.add_paragraph("Baselines were selected to test distinct sources of skill rather than maximize their number:")
bullets(doc, [
    "Seasonal station/basin median and last observation carried forward.",
    "k-nearest neighbors or inverse-distance weighting using geographic distance.",
    "LightGBM/XGBoost using the same static and dynamic covariates.",
    "GRU-D for irregular multivariate events without a river graph.",
    "A directed spatiotemporal graph baseline (DCRNN or Graph WaveNet) on a regularized grid.",
    "Capacity-matched Hydro-CTGSSM variants with no graph, an undirected graph or a Euclidean kNN graph.",
], num_id)
doc.add_paragraph(
    "Required ablations removed, one at a time, continuous-time decay, directed HydroRIVERS edges, edge attributes, auxiliary water-quality states, censored likelihood, GroupDRO and conformal calibration. Hyperparameters were tuned on validation basins only. The final configuration was repeated with five random seeds. Mixed precision and gradient accumulation were used as needed; all configurations, seeds, split manifests and checkpoints were retained."
)

doc.add_heading("2.7 Model evaluation and statistical inference", level=2)
doc.add_paragraph(
    "Regression metrics were MAE, RMSE, R² and Spearman correlation for each target. Probabilistic metrics were negative log-likelihood, continuous ranked probability score where available, 90% interval coverage and mean interval width. Extreme-state detection used AUPRC, AUROC, F1 and event recall; extreme labels were station-season quantiles and were not described as universal regulatory exceedances."
)
doc.add_paragraph(
    "Differences between Hydro-CTGSSM and the strongest baseline were summarized across five seeds and with basin-block bootstrap 95% confidence intervals. Worst-group MAE, group-performance variance, direction-violation diagnostics and errors across sampling-gap, climate, land-use and human-pressure strata were reported. Statistical significance was secondary to effect size, uncertainty and consistency across basin splits."
)

doc.add_heading("2.8 Rapid phosphate-intervention experiment", level=2)
doc.add_heading("2.8.1 Rationale and design", level=3)
doc.add_paragraph(
    "The experiment tested a management response within the water-chemistry envelope observed globally. It did not test river-network propagation and was not presented as a novel-material study. GEMStat river quantiles informed pH 7.0, 7.9 and 8.5 and EC 100, 418 and 1000 µS cm⁻¹. A four-corner plus center design retained boundary information while limiting the main experiment to 30 independent bottles."
)
add_table(doc,
    ["Chemistry group", "pH", "EC (µS cm⁻¹)", "Control bottles", "GFH bottles", "Total"],
    [
        ("B1", "7.0", "100", "3", "3", "6"),
        ("B2", "7.0", "1000", "3", "3", "6"),
        ("B3", "8.5", "100", "3", "3", "6"),
        ("B4", "8.5", "1000", "3", "3", "6"),
        ("Center", "7.9", "418", "3", "3", "6"),
    ],
    [1600, 900, 1800, 1800, 1800, 1460],
    [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER]
)
doc.add_paragraph(
    "Each 250-mL serum bottle contained 200 mL synthetic water with initial DIP 0.40 mg-P L⁻¹, approximately the audited 95th percentile of recent river DIP observations. Controls contained particle-size-matched quartz sand where available; otherwise, no-material blanks were used and this limitation was reported. A single commercial GFH batch was used, with lot, particle size, wet/dry conversion and pretreatment recorded."
)

doc.add_heading("2.8.2 Same-day dose gate", level=3)
doc.add_paragraph(
    "A six-bottle pilot used central chemistry, a no-material blank and two GFH doses (initially 0.10 and 0.50 g L⁻¹), each in duplicate. The lowest dose producing approximately 30–80% removal at 60 min, with blank loss ≤10% and replicate relative standard deviation ≤15%, was selected. If both doses produced >90% or <20% removal, the dose was adjusted before the main run."
)

doc.add_heading("2.8.3 Reaction and analysis", level=3)
doc.add_paragraph(
    "Bottles were maintained at 25±1 °C in the dark and mixed at 150 rpm. Samples (≤2.0 mL) were collected at 0, 10, 30 and 60 min, filtered through validated 0.45-µm membranes and analyzed as soluble reactive phosphorus/DIP by the molybdenum-blue method at 880 nm. Initial and final pH and EC were measured. Dissolved Fe at 60 min and TP were optional risk/QC measurements. MOPS or other organic buffers were not used because they could compete for iron-oxide surface sites."
)
doc.add_paragraph(
    "The calibration curve included at least six standards spanning 0–0.60 mg-P L⁻¹ and required R²≥0.995 without systematic residuals. Method blanks, ≥10% analytical duplicates and matrix spikes were included. Acceptance limits were 80–120% spike recovery and duplicate relative percent difference ≤10%. Bottle positions and sample labels were randomized, and exclusion rules were fixed before unblinding."
)

doc.add_heading("2.8.4 Experimental endpoints and models", level=3)
add_equation(doc, "R₆₀ = 100(C₀ − C₆₀)/C₀", 14)
add_equation(doc, "qobs,60 = (C₀ − C₆₀)V/m", 15)
doc.add_paragraph(
    "C60 was the primary endpoint; R60, qobs,60 and the trapezoidal AUC0–60 were secondary. qobs,60 was described as observed uptake and not maximum adsorption capacity. A prespecified effect model used standardized pH p and log-EC e:"
)
add_equation(doc, "R₆₀ = β₀ + βM M + βp p + βe e + βMp(Mp) + βMe(Me) + βpe(pe) + ε", 16)
doc.add_paragraph(
    "The material indicator M equaled 1 for GFH. HC3 robust confidence intervals and chemistry-cell bootstrap intervals were reported. With n=30, no deep model or large hyperparameter search was permitted. Exploratory early prediction used Ridge regression with C0, C10, C30, pH0, pH30, EC0, EC30 and material status to predict C60. All observations from a bottle remained in one fold, and leave-one-chemistry-group-out validation assessed transfer to an unseen matrix. Elastic Net was a sensitivity analysis; the mean and C30-only models were mandatory baselines."
)

doc.add_heading("2.8.5 Optional natural-water check", level=3)
doc.add_paragraph(
    "If feasible within the second day, one local river-water sample was tested in six bottles (control n=3, GFH n=3) after measuring background DIP/TP, pH, EC and turbidity and spiking DIP to 0.40 mg-P L⁻¹. This was treated as a direction-of-effect reality check, not as broad external validation."
)

doc.add_heading("2.9 Reproducibility and ethics", level=2)
doc.add_paragraph(
    "Data transformations, station matching, split manifests, model configurations and random seeds were version controlled. The open GEMStat data were used under their stated open-data terms with attribution. No human participants or vertebrate animals were involved. Laboratory chemical wastes were handled under institutional safety procedures."
)

doc.add_heading("3. Results", level=1)
doc.add_heading("3.1 Observational structure supports irregular multi-parameter modelling", level=2)
doc.add_paragraph(
    "The audit confirmed that the archive is large but not temporally dense. Among 2010 onward river records, NO3-N comprised 494,069 observations from 12,067 stations, TP 455,413 observations from 9,714 stations, DO 376,504 observations from 10,590 stations and TSS 259,371 observations from 8,343 stations. Across targets, approximately 80–90% of consecutive observations were separated by 8–90 days. Left/right censoring affected 8.72% of NO3-N, 9.10% of TP and 16.49% of TSS records in the full archive."
)
doc.add_paragraph(
    "Multi-parameter event modelling was supported by 344,250 same-station-day NO3-N–TP events and extensive co-observation of temperature, pH, EC and DO. Recent river DIP records numbered 127,337 across 5,167 stations; the median and 95th percentile were 0.046 and 0.378 mg-P L⁻¹, respectively. DIP and TP co-occurred on 84,014 station-days. These distributions justified both the shared-state architecture and the experimental starting concentration."
)
add_caption(doc, "Table 1. Audited coverage of the primary targets in river records from 2010 onward.")
add_table(doc,
    ["Target", "Observations", "Stations", "Countries", "8–31 d (%)", "32–90 d (%)", ">90 d (%)"],
    [
        ("NO3-N", "494,069", "12,067", "30", "43.7", "43.3", "10.2"),
        ("TP", "455,413", "9,714", "26", "43.2", "44.2", "9.4"),
        ("DO", "376,504", "10,590", "27", "41.8", "43.7", "11.4"),
        ("TSS", "259,371", "8,343", "30", "36.8", "48.1", "12.9"),
    ],
    [1400, 1500, 1350, 1100, 1350, 1350, 1310],
)
placeholder(doc, "[INSERT AFTER NETWORK MATCHING: retained event count, retained station count, matched reach count, basin count, match rate, match-distance median/IQR, and excluded-record flow diagram.]")

doc.add_heading("3.2 Whole-basin transfer exposes optimistic random-split performance", level=2)
placeholder(doc, "[TABLE 2 / FIGURE 3 RESULTS: report each model and target under random, forward-time and whole-basin splits; five-seed mean ± SD and basin-block bootstrap 95% CI.]")
doc.add_paragraph(
    "Results-ready text: Under random splitting, Hydro-CTGSSM achieved a mean MAE of [ ] across the four primary targets, compared with [ ] for the strongest baseline. Under whole-basin holdout, errors increased by [ ]%, revealing the optimistic bias of record-level splitting. Hydro-CTGSSM nevertheless reduced whole-basin MAE by [ ]% for NO3-N, [ ]% for TP, [ ]% for DO and [ ]% for TSS relative to [baseline]. The largest transferable gain occurred in [target/condition], whereas [target/condition] showed no material improvement."
)

doc.add_heading("3.3 Direction, event time and censoring make distinct contributions", level=2)
placeholder(doc, "[TABLE 3 / FIGURE 4 RESULTS: ablation deltas with confidence intervals; no graph; Euclidean graph; undirected graph; no decay; regular monthly grid; substitution for censoring; no GroupDRO.]")
doc.add_paragraph(
    "Results-ready text: Replacing the directed HydroRIVERS graph with a Euclidean graph changed whole-basin MAE by [ ] and increased direction-violation diagnostics from [ ] to [ ]. Removing continuous-time decay had little effect for gaps <[ ] days but increased error by [ ]% for gaps >[ ] days. Replacing the censored likelihood with [substitution rule] biased the [lower/upper] tail of [target] and changed extreme-event recall by [ ]. These effects demonstrate [or fail to demonstrate] that each component addresses a distinct monitoring constraint."
)

doc.add_heading("3.4 Calibration reveals model failure boundaries", level=2)
placeholder(doc, "[FIGURE 5 RESULTS: nominal versus empirical coverage, interval width, groupwise coverage, worst-group MAE, and failure maps by climate, basin size, human pressure, gap length and concentration regime.]")
doc.add_paragraph(
    "Results-ready text: Before conformal adjustment, the nominal 90% interval covered [ ]% of whole-basin test observations. Group-aware calibration changed coverage to [ ]% with a mean width of [ ]. Miscoverage remained concentrated in [regions/conditions], defining an operational failure boundary. GroupDRO changed average MAE by [ ] but reduced worst-group MAE by [ ] and between-group variance by [ ]."
)

doc.add_heading("3.5 The rapid experiment tests the management relevance of chemistry boundaries", level=2)
placeholder(doc, "[INSERT QC FIRST: pilot dose; calibration R²; blank loss; spike recovery; duplicate RPD; final n; any preregistered exclusions.]")
placeholder(doc, "[FIGURE 6 RESULTS: C60/R60 by five chemistry groups and material; effect-model coefficients with 95% CI; time courses; early Ridge validation; optional natural-water check.]")
doc.add_paragraph(
    "Results-ready text: At 60 min, GFH changed C60 by [ ] mg-P L⁻¹ (95% CI [ , ]) and R60 by [ ] percentage points. The material×pH interaction was [ ], whereas the material×EC interaction was [ ]. Removal was lowest in [chemistry group], which [did/did not] overlap with the model-derived uncertainty boundary. A Ridge model using observations through 30 min predicted C60 with leave-one-chemistry-group-out MAE [ ], compared with [ ] for the C30-only baseline. In local river water, the material effect was [direction and magnitude], with the stated limitation of one matrix."
)

doc.add_heading("3.6 Reliability-guided monitoring and management priorities", level=2)
placeholder(doc, "[INSERT DECISION ANALYSIS: station/basin priority score, threshold, number/percentage of reaches prioritized, and sensitivity to budget.]")
doc.add_paragraph(
    "Results-ready text: Priority regions were defined by the joint occurrence of high predicted risk, wide calibrated intervals and high expected information gain from a new observation. Compared with uniform allocation, reliability-guided allocation captured [ ]% more high-risk events at the same monitoring budget [or reduced mean interval width by [ ]%]. Management recommendations were stratified into: reliable prediction with effective intervention; reliable prediction with weak intervention; uncertain prediction requiring monitoring; and joint inference/intervention failure requiring site-specific study."
)

doc.add_heading("4. Discussion", level=1)
doc.add_heading("4.1 Cross-basin water-quality learning is an evaluation problem before it is an architecture problem", level=2)
doc.add_paragraph(
    "The principal comparison is between random and whole-basin evaluation. A large random-split score would be scientifically weak if performance collapses when an entire basin is absent from training. The final discussion should therefore lead with the observed transfer gap [INSERT VALUE], not with the most favorable metric. If the directed model retains an advantage under whole-basin holdout, the interpretation is that hydrological connectivity and shared catchment descriptors encode structure that is more transferable than station identity. If the advantage appears only under random splitting, the graph contribution must be downgraded."
)
doc.add_paragraph(
    "The audited sampling intervals explain why regular-grid models may be misleading. Most events are separated by weeks or months, and the timing itself can carry information about monitoring practice. The decay encoder uses elapsed time without inventing daily target values. Its benefit should be interpreted conditionally: a larger gain at long gaps would support the intended mechanism; uniform gains could instead reflect capacity or optimization."
)

doc.add_heading("4.2 Censoring and calibrated uncertainty are central to reliable monitoring", level=2)
doc.add_paragraph(
    "Censoring is substantial for nutrients and especially TSS. Treating a reporting limit as an exact value can compress the lower tail and alter station comparisons. The censored likelihood retains the probability statement actually supplied by the laboratory. However, it assumes that limits and censoring flags are correctly recorded and that the transformed Gaussian family is adequate. Residual diagnostics and a Student-t sensitivity analysis should be used to test, not assume, this adequacy."
)
doc.add_paragraph(
    "Conformal calibration adds an empirical safeguard but does not make every local interval valid. Exchangeability is most questionable during distribution shifts, exactly where monitoring decisions matter. Accordingly, the manuscript reports coverage across climate, basin size, human pressure, gap length and concentration regime. A region with wide intervals may still be responsibly represented; a region with systematic undercoverage is a model failure boundary."
)

doc.add_heading("4.3 Linking predictive failure to intervention boundaries", level=2)
doc.add_paragraph(
    "The bottle experiment serves a narrow but important purpose. It tests whether a management intervention remains effective across a chemistry envelope derived before the experiment. The design avoids claiming that a batch bottle reproduces river transport, biofilm processes or watershed inputs. Instead, it quantifies how pH and EC condition short-contact DIP removal and whether early monitoring can anticipate the endpoint. This is directly relevant to treatment polishing, reactive media deployment or emergency nutrient-control scenarios, while remaining distinct from full-scale river remediation."
)
doc.add_paragraph(
    "A match between model uncertainty and weak GFH performance would identify a joint challenge: the same chemistry regime is difficult both to infer and to manage with the selected intervention. A mismatch is equally informative. Reliable prediction with weak removal suggests that monitoring is adequate but the intervention is unsuitable; uncertain prediction with strong removal suggests that management may be feasible but evidence is insufficient to target it confidently. The four-state framework prevents predictive accuracy from being conflated with treatment effectiveness."
)

doc.add_heading("4.4 Implications for monitoring design and management", level=2)
doc.add_paragraph(
    "Three management outputs should be emphasized after results are available. First, calibrated interval width and empirical miscoverage identify where additional measurements are needed, rather than simply where concentrations are predicted to be high. Second, gap-stratified errors support differentiated sampling frequency: stations whose uncertainty expands rapidly with elapsed time require shorter revisit intervals. Third, whole-basin failure maps can identify hydroclimatic or anthropogenic domains missing from the training archive, guiding network expansion toward representativeness rather than convenience."
)
doc.add_paragraph(
    "The intervention experiment adds an operating-envelope layer. Decisions should distinguish conditions with effective rapid removal from those requiring longer contact, higher dose, alternative media or source control. Because the experiment uses one commercial medium and short contact time, it supports screening and boundary identification, not universal design values."
)

doc.add_heading("4.5 Limitations and boundary of inference", level=2)
bullets(doc, [
    "GEMStat coverage is globally broad but geographically imbalanced and inherits national differences in sampling design, analytical method and reporting limits.",
    "HydroRIVERS matching can be uncertain near confluences, reservoirs, canals and coastal reaches; all network claims depend on the match audit.",
    "Concentration propagation is not mass conservation. Incomplete discharge, unmonitored tributaries and local sources preclude a strong load-balance claim.",
    "GroupDRO performance depends on meaningful group definitions and sufficient data per group; very small basins cannot support stable worst-group estimates.",
    "Conformal intervals provide marginal or empirical groupwise calibration under the chosen split, not universal conditional guarantees under arbitrary future shifts.",
    "The 30-bottle experiment is powered for large boundary effects, not detailed adsorption kinetics or isotherms. One starting concentration cannot estimate qmax.",
    "The optional six-bottle river-water test is a reality check from one matrix and should not be described as independent multi-river validation.",
    "Associations between errors and human-pressure covariates are not causal effects; post hoc feature attribution must not be interpreted as intervention causality.",
], num_id)

doc.add_heading("5. Conclusions", level=1)
placeholder(doc, "[REPLACE WITH 4–5 QUANTIFIED CONCLUSIONS AFTER ALL RESULTS ARE FROZEN.]")
doc.add_paragraph(
    "Planned conclusion structure: (1) quantify the random-to-whole-basin generalization gap; (2) quantify the contribution of direction, elapsed time and censoring; (3) report calibrated coverage and the dominant failure boundary; (4) report the pH/EC-dependent GFH effect and early-prediction performance; and (5) state the monitoring or management action supported by those results. Do not repeat the abstract or claim causality."
)

doc.add_heading("Declarations", level=1)
doc.add_heading("Funding", level=2)
placeholder(doc, "[INSERT FUNDER, GRANT NUMBER AND FUNDER ROLE, OR STATE THAT NO SPECIFIC FUNDING WAS RECEIVED.]")

doc.add_heading("CRediT authorship contribution statement", level=2)
placeholder(doc, "[INSERT AUTHOR-SPECIFIC ROLES: Conceptualization; Methodology; Software; Validation; Formal analysis; Investigation; Data curation; Writing; Visualization; Supervision; Funding acquisition.]")

doc.add_heading("Declaration of competing interest", level=2)
doc.add_paragraph("The authors declare that they have no known competing financial interests or personal relationships that could have appeared to influence the work reported in this paper. [AUTHORS TO CONFIRM]")

doc.add_heading("Declaration of generative AI and AI-assisted technologies", level=2)
doc.add_paragraph(
    "During preparation of this manuscript, the authors used OpenAI Codex to assist with drafting, code organization and document formatting. The authors subsequently reviewed and edited all content and take full responsibility for the scientific design, analysis, interpretation and final text. [ADAPT TO THE JOURNAL POLICY AND ACTUAL USE AT SUBMISSION.]"
)

doc.add_heading("Data and code availability", level=2)
doc.add_paragraph(
    "GEMStat GFQA v3 is available from Zenodo at https://doi.org/10.5281/zenodo.18459694 under the repository terms. HydroRIVERS and HydroATLAS are available from HydroSHEDS. The processed split manifests, station-match audit, model configuration files and analysis code will be archived at [REPOSITORY/DOI] upon acceptance or made available during review [SELECT POLICY]. Raw third-party data will not be redistributed beyond their licenses."
)

doc.add_heading("Acknowledgements", level=2)
placeholder(doc, "[ACKNOWLEDGE DATA PROVIDERS, LABORATORY SUPPORT, COLLABORATORS AND COMPUTING RESOURCES.]")

doc.add_heading("References", level=1)
references = [
    "Carpenter, S.R., Caraco, N.F., Correll, D.L., Howarth, R.W., Sharpley, A.N., Smith, V.H., 1998. Nonpoint pollution of surface waters with phosphorus and nitrogen. Ecological Applications 8, 559–568.",
    "Che, Z., Purushotham, S., Cho, K., Sontag, D., Liu, Y., 2018. Recurrent neural networks for multivariate time series with missing values. Scientific Reports 8, 6085. https://doi.org/10.1038/s41598-018-24271-9.",
    "Chen, T., Guestrin, C., 2016. XGBoost: A scalable tree boosting system. Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining, 785–794.",
    "Heinle, M., Lisniak, D., Saile, P., 2026. UNEP GEMS/Water Global Freshwater Quality Archive, version 3. Zenodo. https://doi.org/10.5281/zenodo.18459694.",
    "Helsel, D.R., 2012. Statistics for Censored Environmental Data Using Minitab and R, 2nd ed. Wiley, Hoboken.",
    "Hilbrandt, I., Lehmann, V., Zietzschmann, F., Ruhl, A.S., Jekel, M., 2019. Quantification and isotherm modelling of competitive phosphate and silicate adsorption onto micro-sized granular ferric hydroxide. RSC Advances 9, 23642–23651. https://doi.org/10.1039/C9RA04865K.",
    "Ke, G., Meng, Q., Finley, T., Wang, T., Chen, W., Ma, W., Ye, Q., Liu, T.-Y., 2017. LightGBM: A highly efficient gradient boosting decision tree. Advances in Neural Information Processing Systems 30.",
    "Lehner, B., Verdin, K., Jarvis, A., 2008. New global hydrography derived from spaceborne elevation data. Eos, Transactions American Geophysical Union 89, 93–94.",
    "Li, Y., Yu, R., Shahabi, C., Liu, Y., 2018. Diffusion convolutional recurrent neural network: Data-driven traffic forecasting. International Conference on Learning Representations.",
    "Linke, S., Lehner, B., Ouellet Dallaire, C., Ariwi, J., Grill, G., Anand, M., Beames, P., Burchard-Levine, V., Maxwell, S., Moidu, H., Tan, F., Thieme, M., 2019. Global hydro-environmental sub-basin and river reach characteristics at high spatial resolution. Scientific Data 6, 283. https://doi.org/10.1038/s41597-019-0300-6.",
    "Romano, Y., Patterson, E., Candès, E.J., 2019. Conformalized quantile regression. Advances in Neural Information Processing Systems 32, 3543–3553.",
    "Sagawa, S., Koh, P.W., Hashimoto, T.B., Liang, P., 2020. Distributionally robust neural networks for group shifts: On the importance of regularization for worst-case generalization. International Conference on Learning Representations.",
    "Smith, V.H., Tilman, G.D., Nekola, J.C., 1999. Eutrophication: Impacts of excess nutrient inputs on freshwater, marine, and terrestrial ecosystems. Environmental Pollution 100, 179–196.",
    "Tobin, J., 1958. Estimation of relationships for limited dependent variables. Econometrica 26, 24–36.",
    "Virro, H., Amatulli, G., Kmoch, A., Shen, L., Uuemaa, E., 2021. GRQA: Global River Water Quality Archive. Earth System Science Data 13, 5483–5507. https://doi.org/10.5194/essd-13-5483-2021.",
    "Vovk, V., Gammerman, A., Shafer, G., 2005. Algorithmic Learning in a Random World. Springer, New York.",
    "Wu, Z., Pan, S., Long, G., Jiang, J., Zhang, C., 2019. Graph WaveNet for deep spatial-temporal graph modeling. Proceedings of the 28th International Joint Conference on Artificial Intelligence, 1907–1913.",
]
for ref in references:
    p = doc.add_paragraph(ref)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)

doc.add_heading("Appendix A. Required result objects before submission", level=1)
add_table(doc,
    ["Object", "Minimum content", "Status"],
    [
        ("Data flow diagram", "Raw → QC → river/2010+ → network-matched → split counts", "Pending"),
        ("Station-match audit", "Match rate, distances, manual sample, fallback decision", "Pending"),
        ("Main performance table", "Four targets × three splits × baselines × five seeds", "Pending GPU"),
        ("Ablation table", "Direction, decay, censoring, covariates, GroupDRO, calibration", "Pending GPU"),
        ("Reliability analysis", "Coverage/width overall and by prespecified groups", "Pending GPU"),
        ("Failure map", "Error, miscoverage and monitoring priority", "Pending GPU"),
        ("Bottle QC", "Dose gate, standards, blanks, spikes, duplicates, exclusions", "Pending lab"),
        ("Bottle effects", "C60/R60, interactions, Ridge group-CV, natural-water check", "Pending lab"),
    ],
    [2200, 5260, 1900],
)

doc.add_heading("Appendix B. Proposed main figures and tables", level=1)
numbered(doc, [
    "Figure 1. Closed-loop study design: GEMStat audit, river network, Hydro-CTGSSM, reliability map and rapid intervention experiment.",
    "Figure 2. Global station distribution, basin splits, sampling intervals, censoring and parameter co-occurrence.",
    "Figure 3. Model performance under random, forward-time and whole-basin holdouts.",
    "Figure 4. Ablation effects across sampling gaps and targets, including directed versus Euclidean graphs.",
    "Figure 5. Calibrated uncertainty, groupwise miscoverage and model failure boundaries.",
    "Figure 6. GFH time courses and R60 response surface across pH/EC boundaries, with early prediction and natural-water check.",
    "Table 1. Audited dataset structure (partially populated in this draft).",
    "Table 2. Baseline and main-model performance with confidence intervals.",
    "Table 3. Ablation, robustness and calibration metrics.",
], num_id)

doc.core_properties.title = "Water Research full manuscript draft: Hydro-CTGSSM and rapid phosphate-intervention experiment"
doc.core_properties.subject = "Results-pending complete manuscript draft"
doc.core_properties.author = "Authors to be confirmed"
doc.core_properties.keywords = "Water Research; GEMStat; HydroRIVERS; graph learning; censored data; phosphate"
doc.core_properties.comments = "Bracketed fields are placeholders and not scientific results."

doc.save(OUT)
print(OUT.resolve())
