from pathlib import Path
import csv
import re
import zipfile
from lxml import etree
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
main_path = ROOT / "outputs" / "Water_Research投稿版_匿名主文稿.docx"
review_path = ROOT / "outputs" / "Water_Research最终稿_审稿与修改记录.docx"
audit_path = ROOT / "outputs" / "Water_Research最终稿_数值单位核查表.csv"
ref_audit_path = ROOT / "outputs" / "Water_Research最终稿_引文核查表.csv"

errors = []
warnings = []


def inspect_doc(path, is_main=False):
    doc = Document(path)
    sec = doc.sections[0]
    if round(sec.page_width.inches, 2) != 8.5 or round(sec.page_height.inches, 2) != 11.0:
        errors.append(f"{path.name}: page size is not Letter")
    for side, value in [("top", sec.top_margin.inches), ("bottom", sec.bottom_margin.inches), ("left", sec.left_margin.inches), ("right", sec.right_margin.inches)]:
        if abs(value - 1.0) > 0.01:
            errors.append(f"{path.name}: {side} margin {value}")
    for name, size in [("Normal", 10.5), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11.5)]:
        style = doc.styles[name]
        if style.font.name != "Arial" or abs(style.font.size.pt - size) > 0.1:
            errors.append(f"{path.name}: style mismatch {name}")

    for ti, table in enumerate(doc.tables, 1):
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        if tbl_w is None or int(tbl_w.get(qn("w:w"))) != 9360:
            errors.append(f"{path.name}: table {ti} width")
        if tbl_ind is None or int(tbl_ind.get(qn("w:w"))) != 120:
            errors.append(f"{path.name}: table {ti} indent")
        grid = [int(x.get(qn("w:w"))) for x in table._tbl.tblGrid]
        if sum(grid) != 9360:
            errors.append(f"{path.name}: table {ti} grid sum {sum(grid)}")
        if not table.rows[0]._tr.get_or_add_trPr().findall(qn("w:tblHeader")):
            errors.append(f"{path.name}: table {ti} no repeating header")
        for ri, row in enumerate(table.rows, 1):
            widths = []
            for cell in row.cells:
                tcw = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
                widths.append(int(tcw.get(qn("w:w"))) if tcw is not None else -1)
            if widths != grid:
                errors.append(f"{path.name}: table {ti} row {ri} cell widths {widths} != {grid}")
                break

    full = "\n".join(p.text for p in doc.paragraphs)
    if "\ufffd" in full or "鈥" in full:
        errors.append(f"{path.name}: broken encoding")
    if re.search(r"\bTODO\b", full):
        errors.append(f"{path.name}: TODO token")
    if is_main:
        headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        for required in ["Abstract", "1. Introduction", "2. Materials and methods", "3. Results", "4. Discussion", "5. Conclusions", "References"]:
            if required not in headings:
                errors.append(f"missing heading {required}")
        if len(doc.inline_shapes) != 5:
            errors.append(f"main: expected 5 figures, found {len(doc.inline_shapes)}")
        for i, shape in enumerate(doc.inline_shapes, 1):
            if not shape._inline.docPr.get("descr"):
                errors.append(f"main: figure {i} lacks alt text")
        ai = headings.index("Abstract") if "Abstract" in headings else -1
        abstract_index = next((i for i, p in enumerate(doc.paragraphs) if p.text == "Abstract"), None)
        if abstract_index is not None:
            words = len(doc.paragraphs[abstract_index + 1].text.split())
            if not 150 <= words <= 300:
                warnings.append(f"abstract word count {words}")
        highlights_index = next((i for i, p in enumerate(doc.paragraphs) if p.text == "Highlights"), None)
        if highlights_index is not None and abstract_index is not None:
            highlights = [p.text for p in doc.paragraphs[highlights_index + 1:abstract_index] if p.text]
            if not 3 <= len(highlights) <= 5:
                errors.append(f"highlight count {len(highlights)}")
            if any(len(x) > 85 for x in highlights):
                errors.append("highlight exceeds 85 characters")
        pending = [p.text for p in doc.paragraphs if "PENDING" in p.text or "[" in p.text and "]" in p.text]
        if not pending:
            errors.append("main: expected explicit pending fields")
        warnings.append(f"main contains {len(pending)} explicit pending/confirmation paragraphs")
    return doc, full


main, main_text = inspect_doc(main_path, True)
review, review_text = inspect_doc(review_path, False)

with ref_audit_path.open(encoding="utf-8-sig") as f:
    refs = list(csv.DictReader(f))
for row in refs:
    if row["cited_in_text"] != "YES":
        errors.append("uncited reference: " + row["citation_key"])

with audit_path.open(encoding="utf-8-sig") as f:
    audit_rows = list(csv.DictReader(f))
if len(audit_rows) < 100:
    errors.append(f"numeric audit unexpectedly short: {len(audit_rows)}")

for path in [main_path, review_path]:
    with zipfile.ZipFile(path) as z:
        xml = etree.fromstring(z.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        if not xml.xpath("//w:lnNumType", namespaces=ns):
            errors.append(f"{path.name}: no line numbering")

print("ERRORS", len(errors))
for x in errors:
    print("ERROR", x)
print("WARNINGS", len(warnings))
for x in warnings:
    print("WARNING", x)
print("MAIN_PARAGRAPHS", len(main.paragraphs), "TABLES", len(main.tables), "FIGURES", len(main.inline_shapes))
print("REVIEW_PARAGRAPHS", len(review.paragraphs), "TABLES", len(review.tables))
print("NUMERIC_AUDIT_ROWS", len(audit_rows))
raise SystemExit(1 if errors else 0)
